#!/usr/bin/env python3
"""Convert Excel workbooks (.xlsx) to Markdown.

Features:
- All worksheets converted to GFM tables
- Formula preservation as backtick-wrapped expressions
- Optional cached formula values
- Formula dependency maps
- Cell formatting/color maps
- Merged cell detection
- Cell comments
- Number/date formatting
- Error recovery per element
"""

from __future__ import annotations

import argparse
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet

from markdownpick.errors import CorruptFileError
from markdownpick.excel.cells import (
    detect_merged_cells,
    excel_col_label,
    get_cell_comment,
    stringify_cell,
)
from markdownpick.excel.formatting import build_formatting_map
from markdownpick.excel.formulas import build_formula_dependency_map
from markdownpick.frontmatter import build_frontmatter
from markdownpick.report import ConversionReport, safe_convert
from markdownpick.tables import markdown_table


def _md_escape(value: str) -> str:
    return value.replace("|", "\\|").replace("\n", "<br>")


def validate_xlsx(path: Path) -> None:
    """Verify a file is a valid OOXML zip package."""
    try:
        with zipfile.ZipFile(path) as zf:
            if "[Content_Types].xml" not in zf.namelist():
                raise CorruptFileError(f"Not a valid OOXML package: {path}")
    except zipfile.BadZipFile as exc:
        raise CorruptFileError(f"File is not a valid zip archive: {path}") from exc


def is_meaningful_cell(cell: Cell) -> bool:
    """Return True when a cell contains meaningful content."""
    if cell.value is None:
        return False
    if cell.data_type == "f":
        return True
    if isinstance(cell.value, str):
        return cell.value.strip() != ""
    return True


def detect_used_bounds(sheet: Worksheet) -> tuple[int, int]:
    """Detect the furthest row/column that contains meaningful content.

    Uses iter_rows() instead of private _cells API.
    """
    max_used_row = 0
    max_used_col = 0

    for row in sheet.iter_rows():
        for cell in row:
            if not is_meaningful_cell(cell):
                continue
            if cell.row > max_used_row:
                max_used_row = cell.row
            if cell.column > max_used_col:
                max_used_col = cell.column

    return max_used_row, max_used_col


def build_markdown_table(
    sheet: Worksheet,
    value_sheet: Worksheet | None,
    include_cell_refs: bool,
    show_formula_values: bool,
    trim_trailing_empty_rows: bool,
    trim_trailing_empty_cols: bool,
    merged_map: dict[str, bool] | None = None,
    show_comments: bool = False,
) -> tuple[str, int, int]:
    """Build a markdown table for one worksheet."""
    used_max_row, used_max_col = detect_used_bounds(sheet)

    max_row = used_max_row if trim_trailing_empty_rows else (sheet.max_row or 1)
    max_col = used_max_col if trim_trailing_empty_cols else (sheet.max_column or 1)

    if max_row <= 0:
        max_row = 1
    if max_col <= 0:
        max_col = 1

    headers: list[str] = []
    if include_cell_refs:
        headers.append("Row")
    headers.extend(excel_col_label(c) for c in range(1, max_col + 1))

    rows: list[list[str]] = []
    for r in range(1, max_row + 1):
        row_values: list[str] = []
        if include_cell_refs:
            row_values.append(str(r))

        for c in range(1, max_col + 1):
            raw_cell = sheet.cell(row=r, column=c)
            value_cell = value_sheet.cell(row=r, column=c) if value_sheet is not None else None
            row_values.append(
                stringify_cell(raw_cell, value_cell, show_formula_values, merged_map, show_comments)
            )

        rows.append(row_values)

    return markdown_table(headers, rows), max_row, max_col


def _collect_sheet_comments(sheet: Worksheet) -> dict[str, str]:
    """Collect all cell comments on a sheet, keyed by coordinate."""
    comments: dict[str, str] = {}
    for row in sheet.iter_rows():
        for cell in row:
            comment = get_cell_comment(cell)
            if comment and cell.coordinate:
                comments[cell.coordinate] = comment
    return comments


def convert_excel_to_markdown(
    excel_path: Path,
    output_path: Path,
    *,
    include_cell_refs: bool = True,
    show_formula_values: bool = False,
    trim_trailing_empty_rows: bool = True,
    trim_trailing_empty_cols: bool = True,
    include_dependency_map: bool = True,
    include_formatting_map: bool = True,
    show_comments: bool = False,
    report: ConversionReport | None = None,
) -> None:
    """Read workbook and write markdown."""
    validate_xlsx(excel_path)

    raw_wb = load_workbook(excel_path, data_only=False)
    value_wb = (
        load_workbook(excel_path, data_only=True)
        if show_formula_values
        else None
    )

    sheet_names = [s.title for s in raw_wb.worksheets]
    features: list[str] = []
    if include_cell_refs:
        features.append("cell references")
    if show_formula_values:
        features.append("cached formula values")
    if include_dependency_map:
        features.append("formula dependency maps")
    if include_formatting_map:
        features.append("color & formatting maps")
    if show_comments:
        features.append("cell comments")

    fm_data: dict[str, str] = {}
    fm_data["sheets"] = ", ".join(sheet_names)

    sections: list[str] = build_frontmatter(
        source=excel_path.name,
        generator="markdownpick/xlsx",
        features=features,
        description=(
            "Machine-readable markdown conversion of an Excel workbook. "
            "Each sheet is rendered as a GitHub-flavored markdown table preserving "
            "formulas as backtick-wrapped expressions. Supplementary maps capture "
            "formula dependencies and cell formatting/color semantics."
        ),
        **fm_data,
    )

    for idx, raw_sheet in enumerate(raw_wb.worksheets):
        value_sheet = value_wb.worksheets[idx] if value_wb is not None else None

        sections.append(f"## Sheet: {_md_escape(raw_sheet.title)}")
        sections.append("")

        merged_map = safe_convert(
            lambda s=raw_sheet: detect_merged_cells(s),
            fallback={},
            report=report,
            label=f"merged cells on {raw_sheet.title}",
        )

        table_md, max_row, max_col = safe_convert(
            lambda rs=raw_sheet, vs=value_sheet, mm=merged_map: build_markdown_table(
                sheet=rs,
                value_sheet=vs,
                include_cell_refs=include_cell_refs,
                show_formula_values=show_formula_values,
                trim_trailing_empty_rows=trim_trailing_empty_rows,
                trim_trailing_empty_cols=trim_trailing_empty_cols,
                merged_map=mm,
                show_comments=show_comments,
            ),
            fallback=("", 1, 1),
            report=report,
            label=f"table on {raw_sheet.title}",
        )
        sections.append(table_md)

        if include_dependency_map:
            sections.append("")
            sections.append("### Formula Dependency Map")
            sections.append("")
            sections.append(
                safe_convert(
                    lambda s=raw_sheet, mr=max_row, mc=max_col: build_formula_dependency_map(s, mr, mc),
                    fallback="_Formula dependency map unavailable._",
                    report=report,
                    label=f"dependency map on {raw_sheet.title}",
                )
            )

        if include_formatting_map:
            sections.append("")
            sections.append("### Color & Formatting Map")
            sections.append("")
            sections.append(
                safe_convert(
                    lambda s=raw_sheet, mr=max_row, mc=max_col: build_formatting_map(s, mr, mc),
                    fallback="_No cell formatting found on this sheet._",
                    report=report,
                    label=f"formatting map on {raw_sheet.title}",
                )
            )
        sections.append("")

        # Sheet comments section
        if show_comments:
            comments = _collect_sheet_comments(raw_sheet)
            if comments:
                sections.append("### Notes")
                sections.append("")
                for coord, text in sorted(comments.items()):
                    sections.append(f"- **{coord}**: {text}")
                sections.append("")

    output_path.write_text("\n".join(sections).rstrip() + "\n", encoding="utf-8")

    if raw_wb:
        raw_wb.close()
    if value_wb:
        value_wb.close()


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convert an Excel workbook to markdown while preserving formulas."
    )
    parser.add_argument("excel_file", help="Path to .xlsx file")
    parser.add_argument(
        "-o", "--output",
        help="Output markdown file path (default: same name with .md)",
        default=None,
    )
    parser.add_argument(
        "--no-cell-refs",
        action="store_true",
        help="Do not include row number reference column.",
    )
    parser.add_argument(
        "--show-formula-values",
        action="store_true",
        help=(
            "When available, include cached formula result as: `=FORMULA` \u21d2 value."
        ),
    )
    parser.add_argument(
        "--keep-empty-rows",
        action="store_true",
        help="Keep trailing completely empty rows instead of trimming them.",
    )
    parser.add_argument(
        "--keep-empty-columns",
        action="store_true",
        help="Keep trailing completely empty columns instead of trimming them.",
    )
    parser.add_argument(
        "--no-dependency-map",
        action="store_true",
        help="Do not include a formula dependency map section per sheet.",
    )
    parser.add_argument(
        "--no-formatting-map",
        action="store_true",
        help="Do not include a formatting map section per sheet.",
    )
    parser.add_argument(
        "--show-comments",
        action="store_true",
        help="Include cell comments/notes.",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv or sys.argv[1:])
    excel_path = Path(args.excel_file).expanduser().resolve()

    if not excel_path.exists():
        print(f"Error: file not found: {excel_path}", file=sys.stderr)
        return 1

    if excel_path.suffix.lower() != ".xlsx":
        print("Error: only .xlsx files are supported.", file=sys.stderr)
        return 1

    output_path = (
        Path(args.output).expanduser().resolve()
        if args.output
        else excel_path.with_suffix(".md")
    )

    report = ConversionReport()
    try:
        convert_excel_to_markdown(
            excel_path=excel_path,
            output_path=output_path,
            include_cell_refs=not args.no_cell_refs,
            show_formula_values=args.show_formula_values,
            trim_trailing_empty_rows=not args.keep_empty_rows,
            trim_trailing_empty_cols=not args.keep_empty_columns,
            include_dependency_map=not args.no_dependency_map,
            include_formatting_map=not args.no_formatting_map,
            show_comments=args.show_comments,
            report=report,
        )
        print(f"Wrote: {output_path}")
    except CorruptFileError as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1
    except Exception as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1

    if report.has_issues:
        print(f"\nConversion completed with {len(report.warnings)} warning(s) and {len(report.errors)} error(s).")
        for w in report.warnings:
            print(f"  Warning: {w}")
        for e in report.errors:
            print(f"  Error: {e}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
