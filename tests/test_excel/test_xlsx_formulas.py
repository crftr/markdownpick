"""Tests for formula parsing."""

from __future__ import annotations

import openpyxl

from markdownpick.excel.formulas import build_formula_dependency_map, extract_formula_references


def test_simple_reference():
    refs = extract_formula_references("=A1")
    assert "A1" in refs


def test_range_reference():
    refs = extract_formula_references("=SUM(A1:B10)")
    assert "A1:B10" in refs


def test_absolute_reference():
    refs = extract_formula_references("=$A$1")
    assert "$A$1" in refs


def test_sheet_reference():
    refs = extract_formula_references("=Sheet1!A1")
    assert "Sheet1!A1" in refs


def test_multi_reference():
    refs = extract_formula_references("=A1+B2+C3")
    assert "A1" in refs
    assert "B2" in refs
    assert "C3" in refs


def test_no_references():
    refs = extract_formula_references("=42")
    assert refs == []


def test_dependency_map(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Value"
    ws["B1"] = "Formula"
    ws["A2"] = 10
    ws["B2"] = "=A2*2"

    result = build_formula_dependency_map(ws, max_row=2, max_col=2)
    assert "B2" in result
    assert "=A2*2" in result
    assert "A2" in result
