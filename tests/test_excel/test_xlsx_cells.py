"""Tests for cell rendering: merged cells, number formats, comments."""

from __future__ import annotations

from datetime import date

import openpyxl

from markdownpick.excel.cells import (
    detect_merged_cells,
    excel_col_label,
    format_cell_value,
    stringify_cell,
)


def test_col_label_a():
    assert excel_col_label(1) == "A"


def test_col_label_z():
    assert excel_col_label(26) == "Z"


def test_col_label_aa():
    assert excel_col_label(27) == "AA"


def test_format_plain_text(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "hello"
    assert format_cell_value(ws["A1"], None, False) == "hello"


def test_format_formula(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "=SUM(B1:B10)"
    result = format_cell_value(ws["A1"], None, False)
    assert result == "`=SUM(B1:B10)`"


def test_format_formula_with_value(tmp_path):
    wb = openpyxl.Workbook()
    ws_raw = wb.active
    ws_raw["A1"] = "=1+1"

    wb2 = openpyxl.Workbook()
    ws_val = wb2.active
    ws_val["A1"] = 2

    result = format_cell_value(ws_raw["A1"], ws_val["A1"], show_formula_values=True)
    assert "\u21d2" in result


def test_format_date(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = date(2026, 5, 2)
    result = format_cell_value(ws["A1"], None, False)
    assert "2026-05-02" in result


def test_merged_cells_detection(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.merge_cells("A1:C1")
    ws["A1"] = "merged"

    merged = detect_merged_cells(ws)
    assert "A1" in merged
    assert merged["A1"] is True
    assert "B1" in merged
    assert merged["B1"] is False


def test_stringify_merged_non_top(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.merge_cells("A1:C1")
    ws["A1"] = "merged"

    merged = detect_merged_cells(ws)
    result = stringify_cell(ws["B1"], None, False, merged)
    assert result == ""


def test_stringify_merged_top(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.merge_cells("A1:C1")
    ws["A1"] = "merged"

    merged = detect_merged_cells(ws)
    result = stringify_cell(ws["A1"], None, False, merged)
    assert "\u2197" in result


def test_empty_cell(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = None
    assert format_cell_value(ws["A1"], None, False) == ""
