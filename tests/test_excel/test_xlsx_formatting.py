"""Tests for cell formatting maps."""

from __future__ import annotations

import openpyxl

from markdownpick.excel.formatting import (
    _cell_format_descriptor,
    _collapse_ranges,
    build_formatting_map,
)


def test_collapse_ranges_contiguous():
    result = _collapse_ranges(["B6", "B7", "B8"])
    assert result == "B6:B8"


def test_collapse_ranges_gap():
    result = _collapse_ranges(["B6", "B7", "B8", "B13"])
    assert result == "B6:B8, B13"


def test_collapse_ranges_multiple_columns():
    result = _collapse_ranges(["B6", "B7", "D6"])
    assert "B6:B7" in result
    assert "D6" in result


def test_collapse_ranges_single():
    result = _collapse_ranges(["A1"])
    assert result == "A1"


def test_format_descriptor_bold():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "test"
    ws["A1"].font = openpyxl.styles.Font(bold=True)
    desc = _cell_format_descriptor(ws["A1"])
    assert "bold" in desc


def test_format_descriptor_bg_color():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "test"
    ws["A1"].fill = openpyxl.styles.PatternFill(
        start_color="FFFF00", end_color="FFFF00", fill_type="solid"
    )
    desc = _cell_format_descriptor(ws["A1"])
    assert "background" in desc.lower()


def test_format_descriptor_none_for_plain():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "plain"
    desc = _cell_format_descriptor(ws["A1"])
    assert desc is None


def test_build_formatting_map():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "bold"
    ws["A1"].font = openpyxl.styles.Font(bold=True)
    ws["A2"] = "plain"

    result = build_formatting_map(ws, max_row=2, max_col=1)
    assert "bold" in result
