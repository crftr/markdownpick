"""Tests for the docx converter."""

from __future__ import annotations

from pathlib import Path

import docx

from markdownpick.docx.converter import (
    _extract_footnotes,
    convert_docx_to_markdown,
    validate_docx,
)
from markdownpick.errors import CorruptFileError
from markdownpick.report import ConversionReport


def test_convert_styled_docx(docx_styled: Path, tmp_path: Path):
    out = tmp_path / "styled.md"
    convert_docx_to_markdown(docx_styled, out)
    text = out.read_text()
    assert "---" in text
    assert "generator: markdownpick/docx" in text
    assert "# Document Title" in text
    assert "## First Section" in text
    assert "**bold text**" in text
    assert "*italic text*" in text
    assert "***bold italic***" in text
    assert "> A quoted block" in text
    assert "| Name" in text


def test_convert_empty_docx(docx_empty: Path, tmp_path: Path):
    out = tmp_path / "empty.md"
    convert_docx_to_markdown(docx_empty, out)
    text = out.read_text()
    assert "---" in text


def test_convert_underline_sup(docx_with_underline_sup: Path, tmp_path: Path):
    out = tmp_path / "formatting.md"
    convert_docx_to_markdown(docx_with_underline_sup, out)
    text = out.read_text()
    assert "<u>underlined</u>" in text
    assert "<sup>superscript</sup>" in text
    assert "<sub>subscript</sub>" in text


def test_validate_corrupt(corrupt_file: Path):
    try:
        validate_docx(corrupt_file)
    except CorruptFileError:
        return
    raise AssertionError("Should have raised CorruptFileError")


def test_convert_with_report(docx_styled: Path, tmp_path: Path):
    out = tmp_path / "report.md"
    report = ConversionReport()
    convert_docx_to_markdown(docx_styled, out, report=report)
    assert out.exists()


def test_extract_footnotes(tmp_path: Path):
    d = docx.Document()
    d.add_paragraph("Text with footnote[^1].")
    # python-docx doesn't easily create real footnotes via the high-level API,
    # but _extract_footnotes should handle the case where there are none
    path = tmp_path / "no_notes.docx"
    d.save(path)

    loaded = docx.Document(str(path))
    result = _extract_footnotes(loaded)
    assert result == ""
