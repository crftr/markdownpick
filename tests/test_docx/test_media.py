"""Tests for docx media/image extraction."""

from __future__ import annotations

from pathlib import Path

import docx

from markdownpick.docx.media import extract_images, images_to_markdown


def _make_docx_with_image(tmp_path: Path) -> Path:
    """Create a docx with an embedded image (minimal pixel PNG)."""
    d = docx.Document()
    d.add_paragraph("Before image")

    # Add a minimal 1x1 pixel PNG
    png_data = (
        b"\x89PNG\r\n\x1a\n"
        b"\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x02\x00\x00\x00\x90wS\xde"
        b"\x00\x00\x00\x0cIDATx\x9cc\xf8\x0f\x00\x00\x01\x01\x00\x05\x18\xd8N"
        b"\x00\x00\x00\x00IEND\xaeB`\x82"
    )

    try:
        from docx.shared import Inches

        r = d.add_paragraph().add_run()
        # Try to add image via the internal API
        image_part = d.part.new_image_part(png_data)
        r._r._insert_drawing_for_inline(image_part, Inches(1), Inches(1))
    except Exception:
        # If image insertion fails at the API level, skip — this tests the extractor only
        pass

    d.add_paragraph("After image")
    path = tmp_path / "with_image.docx"
    d.save(path)
    return path


def test_extract_images_no_images(tmp_path: Path):
    d = docx.Document()
    d.add_paragraph("No images here")
    path = tmp_path / "no_images.docx"
    d.save(path)

    images_dir = tmp_path / "images"
    result = extract_images(d, images_dir)
    assert result == []


def test_images_to_markdown():
    extracted = [(1, "images/img_001.png", "Test Image")]
    md = images_to_markdown(extracted)
    assert md == [(1, "![Test Image](images/img_001.png)")]
