"""Tests for docx inline formatting."""

from __future__ import annotations

import docx

from markdownpick.docx.inlines import _run_to_inline


def _make_run(text: str, bold=False, italic=False, strike=False, underline=False, sup=False, sub=False):
    d = docx.Document()
    p = d.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.strike = strike
    run.underline = underline
    if sup:
        run.font.superscript = True
    if sub:
        run.font.subscript = True
    return run


def test_plain_text():
    r = _make_run("hello")
    assert _run_to_inline(r) == "hello"


def test_bold():
    r = _make_run("bold", bold=True)
    assert _run_to_inline(r) == "**bold**"


def test_italic():
    r = _make_run("italic", italic=True)
    assert _run_to_inline(r) == "*italic*"


def test_bold_italic():
    r = _make_run("both", bold=True, italic=True)
    assert _run_to_inline(r) == "***both***"


def test_strikethrough():
    r = _make_run("deleted", strike=True)
    assert _run_to_inline(r) == "~~deleted~~"


def test_underline():
    r = _make_run("underline", underline=True)
    assert _run_to_inline(r) == "<u>underline</u>"


def test_superscript():
    r = _make_run("sup", sup=True)
    assert _run_to_inline(r) == "<sup>sup</sup>"


def test_subscript():
    r = _make_run("sub", sub=True)
    assert _run_to_inline(r) == "<sub>sub</sub>"


def test_empty_run():
    r = _make_run("")
    assert _run_to_inline(r) == ""


def test_whitespace_preservation():
    r = _make_run("  bold  ", bold=True)
    result = _run_to_inline(r)
    assert result.startswith("  ")
    assert result.endswith("  ")
    assert "**bold**" in result
