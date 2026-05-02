"""Shared fixtures for markdownpick tests.

All fixtures are created programmatically — no binary files stored in the repo.
"""

from __future__ import annotations

from pathlib import Path

import docx
import openpyxl
import pytest


@pytest.fixture
def tmp_dir(tmp_path: Path) -> Path:
    return tmp_path


# ── docx fixtures ────────────────────────────────────────────────────────────


@pytest.fixture
def docx_styled(tmp_path: Path) -> Path:
    """A docx with headings, styled paragraphs, bold/italic, and a table."""
    d = docx.Document()
    d.add_heading("Document Title", level=1)
    d.add_heading("First Section", level=2)

    p = d.add_paragraph()
    p.add_run("This is ")
    p.add_run("bold text").bold = True
    p.add_run(" and ")
    p.add_run("italic text").italic = True
    p.add_run(" and ")
    r = p.add_run("bold italic")
    r.bold = True
    r.italic = True

    d.add_paragraph("Normal paragraph text.", style="Normal")

    # List items
    d.add_paragraph("Bullet one", style="List Bullet")
    d.add_paragraph("Bullet two", style="List Bullet")
    d.add_paragraph("Number one", style="List Number")

    # Quote
    d.add_paragraph("A quoted block of text.", style="Quote")

    # Table
    table = d.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(0, 2).text = "Notes"
    table.cell(1, 0).text = "Alice"
    table.cell(1, 1).text = "100"
    table.cell(1, 2).text = "Good"

    path = tmp_path / "styled.docx"
    d.save(path)
    return path


@pytest.fixture
def docx_empty(tmp_path: Path) -> Path:
    """A minimal docx with no content."""
    d = docx.Document()
    path = tmp_path / "empty.docx"
    d.save(path)
    return path


@pytest.fixture
def docx_with_underline_sup(tmp_path: Path) -> Path:
    """A docx with underline, superscript, and subscript runs."""
    d = docx.Document()
    p = d.add_paragraph()
    p.add_run("normal ")
    u = p.add_run("underlined")
    u.underline = True
    p.add_run(" ")
    s = p.add_run("superscript")
    s.font.superscript = True
    p.add_run(" ")
    sb = p.add_run("subscript")
    sb.font.subscript = True

    path = tmp_path / "formatting.docx"
    d.save(path)
    return path


@pytest.fixture
def corrupt_file(tmp_path: Path) -> Path:
    """A file that is not a valid zip."""
    path = tmp_path / "corrupt.docx"
    path.write_bytes(b"not a zip file at all")
    return path


# ── xlsx fixtures ────────────────────────────────────────────────────────────


@pytest.fixture
def xlsx_formulas(tmp_path: Path) -> Path:
    """An xlsx with formulas, merged cells, and formatting."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"

    ws["A1"] = "Value"
    ws["B1"] = "Doubled"
    ws["C1"] = "Sum"
    ws["A2"] = 10
    ws["B2"] = "=A2*2"
    ws["C2"] = "=A2+B2"
    ws["A3"] = 20
    ws["B3"] = "=A3*2"
    ws["C3"] = "=A3+B3"

    # Merged cell
    ws.merge_cells("A5:C5")
    ws["A5"] = "This is merged"

    # Formatting
    from openpyxl.styles import Font, PatternFill

    ws["A1"].font = Font(bold=True, color="FF0000")
    ws["A2"].fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

    path = tmp_path / "formulas.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def xlsx_multi_sheet(tmp_path: Path) -> Path:
    """An xlsx with multiple sheets."""
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1["A1"] = "One"
    ws1["B1"] = "Two"

    ws2 = wb.create_sheet("Sheet2")
    ws2["A1"] = "Alpha"
    ws2["A2"] = "Beta"

    path = tmp_path / "multi.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def xlsx_empty(tmp_path: Path) -> Path:
    """An xlsx with no data."""
    wb = openpyxl.Workbook()
    path = tmp_path / "empty.xlsx"
    wb.save(path)
    return path
